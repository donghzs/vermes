# -*- coding: utf-8 -*-
"""A2 选区编辑（局部重生成）后端真实往返测试。

覆盖：line_range 选区 / text_fingerprint 选区 / docx block_index 段落回填 /
非法请求（缺 instruction→400、指纹未命中→404、未注入 regenerator→501 诚实报错）。

注：LLM 重生成步骤用确定性 monkeypatch 替身验证「锚点定位 + patch 写回 + A1 自动快照」全链路；
生产真实 regenerator 由 chat 蓝图启动时注入（见 artifacts.register_region_regenerator）。
"""
import sys
import os
import uuid

sys.path.insert(0, "/Users/dongzusheng/Projects/vermes-electron")

import pytest


@pytest.fixture
def ver_db(tmp_path, monkeypatch):
    """把版本库指向临时文件，避免污染真实 ~/.vermes/artifacts_versions.db。"""
    import vermes_cli.blueprints.artifacts as art
    db = tmp_path / "versions_test.db"
    monkeypatch.setattr(art, "_VERSIONS_DB_PATH", str(db))
    art._init_versions_db()
    yield art
    for ext in ("", "-wal", "-shm"):
        p = str(db) + ext
        if os.path.exists(p):
            os.unlink(p)


@pytest.fixture
def client(ver_db):
    import vermes_cli.web_server as ws
    from starlette.testclient import TestClient
    return TestClient(ws.app)


def _mk_artifact(ext: str) -> str:
    p = f"/tmp/a2p_{uuid.uuid4().hex}{ext}"
    open(p, "w").close()
    return p


def _rel(p: str) -> str:
    return "tmp/" + os.path.basename(p)


def _deterministic_regenerator(transform):
    """返回 (selection, instruction, before, after, model_hint)->text 的确定性替身。"""
    def _fn(selection, instruction, before, after, model_hint=None):
        return transform(selection)
    return _fn


def test_patch_line_range(client, ver_db, monkeypatch):
    """line_range 选区：第 2 行被 regenerator 改写，文件其余不变，且自动记一版。"""
    monkeypatch.setattr(ver_db, "_REGION_REGENERATOR",
                        _deterministic_regenerator(lambda s: s.upper()))
    p = _mk_artifact(".md")
    try:
        open(p, "w").write("line1\nline2\nline3\n")
        rel = _rel(p)
        r = client.post(f"/api/v1/artifacts/{rel}/patch", json={
            "anchor": {"type": "line_range", "start": 2, "end": 2},
            "selection": "line2", "instruction": "改成大写",
        })
        assert r.status_code == 200, r.text
        j = r.json()
        assert j["type"] == "text"
        assert j["before"] == "line2"
        assert j["after"] == "LINE2"
        assert j["line_range"] == [2, 2]
        assert open(p, "r").read() == "line1\nLINE2\nline3\n"
        # A1 自动快照：选区重生成产生一条版本
        v = client.get(f"/api/v1/artifacts/{rel}/versions").json()
        assert len(v["versions"]) == 1
        assert v["versions"][0]["note"] == "选区重生成"
    finally:
        os.unlink(p)


def test_patch_text_fingerprint(client, ver_db, monkeypatch):
    """text_fingerprint 选区：按原文指纹定位并替换（多行选区也能定位）。"""
    monkeypatch.setattr(ver_db, "_REGION_REGENERATOR",
                        _deterministic_regenerator(lambda s: "[" + s + "]"))
    p = _mk_artifact(".md")
    try:
        open(p, "w").write("alpha\nbeta\ngamma\n")
        rel = _rel(p)
        r = client.post(f"/api/v1/artifacts/{rel}/patch", json={
            "anchor": {"type": "text_fingerprint", "text": "beta\ngamma"},
            "instruction": "加方括号",
        })
        assert r.status_code == 200, r.text
        assert open(p, "r").read() == "alpha\n[beta\ngamma]\n"
        assert r.json()["line_range"] == [2, 3]
    finally:
        os.unlink(p)


def test_patch_docx_block_index(client, ver_db, monkeypatch):
    """docx block_index 段落回填：第 1 段被改写，其余段落不动，自动记一版。"""
    monkeypatch.setattr(ver_db, "_REGION_REGENERATOR",
                        _deterministic_regenerator(lambda s: "REGEN:" + s))
    from docx import Document
    p = f"/tmp/a2p_{uuid.uuid4().hex}.docx"
    try:
        d = Document()
        for t in ("A", "B", "C"):
            d.add_paragraph(t)
        d.save(p)
        rel = _rel(p)
        r = client.post(f"/api/v1/artifacts/{rel}/patch", json={
            "anchor": {"type": "block_index", "index": 1},
            "selection": "B", "instruction": "改写这段",
        })
        assert r.status_code == 200, r.text
        assert r.json()["block_index"] == 1
        d2 = Document(p)
        assert [para.text for para in d2.paragraphs] == ["A", "REGEN:B", "C"]
        v = client.get(f"/api/v1/artifacts/{rel}/versions").json()
        assert len(v["versions"]) == 1
        assert "docx" in v["versions"][0]["note"]
    finally:
        if os.path.exists(p):
            os.unlink(p)


def test_patch_missing_instruction_400(client):
    """缺 instruction → 400。"""
    p = _mk_artifact(".md")
    try:
        open(p, "w").write("x\n")
        rel = _rel(p)
        r = client.post(f"/api/v1/artifacts/{rel}/patch", json={
            "anchor": {"type": "line_range", "start": 1, "end": 1}, "selection": "x",
        })
        assert r.status_code == 400
    finally:
        os.unlink(p)


def test_patch_no_regenerator_501(client, ver_db, monkeypatch):
    """未注入真实 regenerator → 501（诚实报错，不静默假成功）。
    显式把接缝置 None，模拟「未配置」场景（生产环境 chat 蓝图会在 import 时注入真实实现）。
    """
    monkeypatch.setattr(ver_db, "_REGION_REGENERATOR", None)
    p = _mk_artifact(".md")
    try:
        open(p, "w").write("x\n")
        rel = _rel(p)
        r = client.post(f"/api/v1/artifacts/{rel}/patch", json={
            "anchor": {"type": "line_range", "start": 1, "end": 1},
            "selection": "x", "instruction": "改写",
        })
        assert r.status_code == 501
        assert "未接入模型" in r.json()["detail"]
    finally:
        os.unlink(p)


def test_patch_fingerprint_not_found_404(client, ver_db, monkeypatch):
    """指纹在当前文件中不存在 → 404。"""
    monkeypatch.setattr(ver_db, "_REGION_REGENERATOR",
                        _deterministic_regenerator(lambda s: s))
    p = _mk_artifact(".md")
    try:
        open(p, "w").write("hello\n")
        rel = _rel(p)
        r = client.post(f"/api/v1/artifacts/{rel}/patch", json={
            "anchor": {"type": "text_fingerprint", "text": "not present"},
            "instruction": "改写",
        })
        assert r.status_code == 404
    finally:
        os.unlink(p)


def test_real_regenerator_registered_by_chat_blueprint():
    """chat 蓝图 import 时应已把真实 regenerator 注入 artifacts 接缝（使 /patch 端到端可用）。"""
    import vermes_cli.blueprints.chat as chat_bp
    import vermes_cli.blueprints.artifacts as art
    assert art._REGION_REGENERATOR is not None
    assert art._REGION_REGENERATOR is chat_bp._region_regenerate
